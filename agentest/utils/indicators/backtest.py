"""Historical backtesting of EMA 10/20 + Heikin-Ashi signals."""

from __future__ import annotations
from collections import defaultdict
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any
import pandas as pd
import yfinance as yf
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from agentest.utils.indicators.ema import compute_ema, detect_crossover
from agentest.utils.indicators.heikin_ashi import compute_heikin_ashi, detect_ha_signal, scan_combined

DOCUMENTS_DIR = Path.home() / "Documents" / "Agentest_Reports"

BT_PARAMS = {
    "lookback_months": 18,
    "min_data_days": 60,
}

HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)
ACTION_FILLS = {
    "BUY":  PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
    "SELL": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
}


def _fetch_data(symbol: str) -> pd.DataFrame | None:
    try:
        t = yf.Ticker(symbol)
        df = t.history(period=f"{BT_PARAMS['lookback_months']}mo", interval="1d")
        if df.empty or len(df) < BT_PARAMS["min_data_days"]:
            return None
        return df
    except Exception:
        return None


def _signal_at_idx(df: pd.DataFrame, idx: int) -> dict:
    sub = df.iloc[: idx + 1]
    if len(sub) < 30:
        return {"action": "SKIP", "reason": "insufficient_data"}
    short = compute_ema(sub["Close"], 10)
    long_ = compute_ema(sub["Close"], 20)
    cross = detect_crossover(short, long_)
    ha = compute_heikin_ashi(sub)
    ha_sig = detect_ha_signal(ha)
    combined = scan_combined(short, long_, ha_sig)
    return combined


def _estimate_scan_freq() -> int:
    """Trading days between scans. 1 = daily, 5 = weekly."""
    return 1


def backtest_symbol(symbol: str, category: str) -> list[dict]:
    """Long-only backtest: BUY → enter, SELL → exit, ignore SELL when flat."""
    df = _fetch_data(symbol)
    if df is None:
        return []
    step = _estimate_scan_freq()
    trades: list[dict] = []
    in_position = False
    entry_price = 0.0
    entry_idx = 0
    entry_date: pd.Timestamp | None = None
    entry_signal = ""

    for i in range(60, len(df), step):
        sig = _signal_at_idx(df, i)
        action = sig.get("action", "SKIP")

        if action == "BUY" and not in_position:
            in_position = True
            entry_price = float(df["Close"].iloc[i])
            entry_idx = i
            entry_date = df.index[i]
            entry_signal = sig.get("reason", "")

        elif action == "SELL" and in_position:
            exit_price = float(df["Close"].iloc[i])
            exit_date = df.index[i]
            pnl_pct = (exit_price - entry_price) / entry_price * 100
            trades.append(dict(
                symbol=symbol, category=category,
                direction="BUY", entry_date=str(entry_date.date()),
                exit_date=str(exit_date.date()),
                entry_price=round(entry_price, 2),
                exit_price=round(exit_price, 2),
                pnl_pct=round(pnl_pct, 2),
                reason=entry_signal,
                bars_held=i - entry_idx,
            ))
            in_position = False

    # Force-close any open position at series end
    if in_position:
        exit_price = float(df["Close"].iloc[-1])
        pnl_pct = (exit_price - entry_price) / entry_price * 100
        trades.append(dict(
            symbol=symbol, category=category,
            direction="BUY", entry_date=str(entry_date.date()),
            exit_date=str(df.index[-1].date()),
            entry_price=round(entry_price, 2),
            exit_price=round(exit_price, 2),
            pnl_pct=round(pnl_pct, 2),
            reason=entry_signal,
            bars_held=len(df) - entry_idx,
        ))

    return trades


def run_backtest(symbols: dict[str, list[str]]) -> dict[str, Any]:
    all_trades: list[dict] = []
    by_category: dict[str, list[dict]] = defaultdict(list)
    perf_by_symbol: dict[str, list[float]] = defaultdict(list)

    for cat, sym_list in symbols.items():
        for sym in sym_list:
            trades = backtest_symbol(sym, cat)
            for t in trades:
                all_trades.append(t)
                by_category[cat].append(t)
                perf_by_symbol[t["symbol"]].append(t["pnl_pct"])

    results: dict[str, Any] = {"trades": all_trades, "categories": {}}

    for cat, trades_list in by_category.items():
        wins = [t for t in trades_list if t["pnl_pct"] > 0]
        losses = [t for t in trades_list if t["pnl_pct"] <= 0]
        total = len(trades_list)
        win_rate = round(len(wins) / total * 100, 1) if total else 0
        avg_pnl = round(sum(t["pnl_pct"] for t in trades_list) / total, 2) if total else 0
        sum_pnl = round(sum(t["pnl_pct"] for t in trades_list), 2)
        gross_profit = sum(t["pnl_pct"] for t in wins) if wins else 0
        gross_loss = abs(sum(t["pnl_pct"] for t in losses)) if losses else 0
        profit_factor = round(gross_profit / gross_loss, 2) if gross_loss else float("inf")
        results["categories"][cat] = dict(
            total_trades=total, wins=len(wins), win_rate=win_rate,
            avg_pnl=avg_pnl, sum_pnl=sum_pnl, profit_factor=profit_factor,
        )

    sym_perf = []
    for sym, pnls in perf_by_symbol.items():
        compound = 1.0
        wins = sum(1 for p in pnls if p > 0)
        for p in pnls:
            compound *= 1 + p / 100
        compound_return = round((compound - 1) * 100, 2)
        sym_perf.append(dict(
            symbol=sym, trades=len(pnls), wins=wins,
            win_rate=round(wins / len(pnls) * 100, 1) if pnls else 0,
            avg_pnl=round(sum(pnls) / len(pnls), 2),
            cum_pnl=round(sum(pnls), 2),
            compound_return=compound_return,
        ))
    sym_perf.sort(key=lambda x: x["compound_return"], reverse=True)
    results["top_performers"] = sym_perf[:20]
    results["worst_performers"] = sym_perf[-20:]
    results["symbol_stats"] = sym_perf  # full list for company-wise output

    total_trades = len(all_trades)
    total_wins = len([t for t in all_trades if t["pnl_pct"] > 0])
    total_losses = [t for t in all_trades if t["pnl_pct"] <= 0]
    gross_profit = sum(t["pnl_pct"] for t in all_trades if t["pnl_pct"] > 0)
    gross_loss = abs(sum(t["pnl_pct"] for t in all_trades if t["pnl_pct"] <= 0))
    results["overall"] = dict(
        total_trades=total_trades, wins=total_wins,
        win_rate=round(total_wins / total_trades * 100, 1) if total_trades else 0,
        avg_pnl=round(sum(t["pnl_pct"] for t in all_trades) / total_trades, 2) if total_trades else 0,
        sum_pnl=round(sum(t["pnl_pct"] for t in all_trades), 2),
        profit_factor=round(gross_profit / gross_loss, 2) if gross_loss else float("inf"),
    )
    return results


def _write_cell(ws, row, col, value, font=None, fill=None):
    cell = ws.cell(row=row, column=col, value=value)
    if font:
        cell.font = font
    if fill:
        cell.fill = fill
    cell.border = THIN_BORDER
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    return cell


def _write_header(ws, row, headers):
    for col, h in enumerate(headers, 1):
        _write_cell(ws, row, col, h, font=HEADER_FONT, fill=HEADER_FILL)
    ws.row_dimensions[row].height = 22


def export_results(results: dict[str, Any]) -> Path:
    ts = datetime.now()
    folder = DOCUMENTS_DIR / ts.strftime("%Y%m%d_%H%M%S")
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"Backtest_{ts.strftime('%Y%m%d_%H%M%S')}.xlsx"

    wb = Workbook()

    # ── Summary sheet ──
    ws = wb.active
    ws.title = "Summary"
    ws.cell(row=1, column=1, value="Backtest Results — EMA 10/20 + Heikin-Ashi").font = Font(bold=True, size=14)
    ws.cell(row=2, column=1, value=f"Generated: {ts:%Y-%m-%d %H:%M:%S}").font = Font(size=10, italic=True)
    ws.merge_cells("A1:G1")

    o = results["overall"]
    row = 4
    for label, val in [("Total Trades", o["total_trades"]), ("Wins", o["wins"]),
                       ("Win Rate", f"{o['win_rate']}%"), ("Avg P&L per Trade", f"{o['avg_pnl']:.2f}%"),
                       ("Total Edge (sum P&L)", f"{o['sum_pnl']:.2f}%"),
                       ("Profit Factor", f"{o['profit_factor']}")]:
        _write_cell(ws, row, 1, label, font=Font(bold=True))
        _write_cell(ws, row, 2, val)
        row += 1

    row += 1
    _write_header(ws, row, ["Category", "Trades", "Wins", "Win Rate", "Avg P&L", "Total Edge", "Profit Factor"])
    row += 1
    for cat in sorted(results["categories"]):
        c = results["categories"][cat]
        fill = ACTION_FILLS.get("BUY") if c["sum_pnl"] > 0 else ACTION_FILLS.get("SELL")
        _write_cell(ws, row, 1, cat, fill=fill)
        _write_cell(ws, row, 2, c["total_trades"], fill=fill)
        _write_cell(ws, row, 3, c["wins"], fill=fill)
        _write_cell(ws, row, 4, f"{c['win_rate']}%", fill=fill)
        _write_cell(ws, row, 5, f"{c['avg_pnl']:.2f}%", fill=fill)
        _write_cell(ws, row, 6, f"{c['sum_pnl']:.2f}%", fill=fill)
        _write_cell(ws, row, 7, c["profit_factor"], fill=fill)
        row += 1
    for c in range(1, 8):
        ws.column_dimensions[get_column_letter(c)].width = 16
    for c in range(1, 7):
        ws.column_dimensions[get_column_letter(c)].width = 16

    # ── Top 20 sheet ──
    ws2 = wb.create_sheet("Top 20")
    _write_header(ws2, 1, ["#", "Symbol", "Trades", "Avg P&L", "Total Edge", "Compound Return"])
    for i, p in enumerate(results["top_performers"], 1):
        fill = ACTION_FILLS.get("BUY") if p["compound_return"] > 0 else ACTION_FILLS.get("SELL")
        _write_cell(ws2, i + 1, 1, i, fill=fill)
        _write_cell(ws2, i + 1, 2, p["symbol"], fill=fill)
        _write_cell(ws2, i + 1, 3, p["trades"], fill=fill)
        _write_cell(ws2, i + 1, 4, f"{p['avg_pnl']:.2f}%", fill=fill)
        _write_cell(ws2, i + 1, 5, f"{p['cum_pnl']:.2f}%", fill=fill)
        _write_cell(ws2, i + 1, 6, f"{p['compound_return']:.2f}%", fill=fill)

    # ── Bottom 20 sheet ──
    ws3 = wb.create_sheet("Bottom 20")
    _write_header(ws3, 1, ["#", "Symbol", "Trades", "Avg P&L", "Total Edge", "Compound Return"])
    for i, p in enumerate(results["worst_performers"], 1):
        fill = ACTION_FILLS.get("SELL") if p["compound_return"] < 0 else ACTION_FILLS.get("BUY")
        _write_cell(ws3, i + 1, 1, i, fill=fill)
        _write_cell(ws3, i + 1, 2, p["symbol"], fill=fill)
        _write_cell(ws3, i + 1, 3, p["trades"], fill=fill)
        _write_cell(ws3, i + 1, 4, f"{p['avg_pnl']:.2f}%", fill=fill)
        _write_cell(ws3, i + 1, 5, f"{p['cum_pnl']:.2f}%", fill=fill)
        _write_cell(ws3, i + 1, 6, f"{p['compound_return']:.2f}%", fill=fill)

    # ── All Trades sheet ──
    ws4 = wb.create_sheet("All Trades")
    _write_header(ws4, 1, ["Symbol", "Category", "Direction", "Entry Date", "Exit Date",
                            "Entry Price", "Exit Price", "P&L %", "Reason", "Bars Held"])
    for i, t in enumerate(results["trades"], 2):
        fill = ACTION_FILLS.get("BUY") if t["pnl_pct"] > 0 else ACTION_FILLS.get("SELL")
        _write_cell(ws4, i, 1, t["symbol"], fill=fill)
        _write_cell(ws4, i, 2, t["category"], fill=fill)
        _write_cell(ws4, i, 3, t["direction"], fill=fill)
        _write_cell(ws4, i, 4, t["entry_date"], fill=fill)
        _write_cell(ws4, i, 5, t["exit_date"], fill=fill)
        _write_cell(ws4, i, 6, t["entry_price"], fill=fill)
        _write_cell(ws4, i, 7, t["exit_price"], fill=fill)
        _write_cell(ws4, i, 8, f"{t['pnl_pct']:.2f}%", fill=fill)
        _write_cell(ws4, i, 9, t["reason"], fill=fill)
        _write_cell(ws4, i, 10, t["bars_held"], fill=fill)

    wb.save(path)
    return path
