"""Enriched Excel report — technical signals + fundamental analysis."""

from __future__ import annotations
import json
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from agentest.utils.indicators.fundamental_analysis import analyze_batch

SCAN_RESULTS_DIR = Path("tests/temp/scan_results")


def _load_all_results() -> list[dict]:
    """Load ALL scan result JSONs (BUY, SELL, HOLD, WATCH, SKIP)."""
    stocks: list[dict] = []
    if not SCAN_RESULTS_DIR.exists():
        return stocks
    for f in sorted(SCAN_RESULTS_DIR.iterdir()):
        if f.suffix == ".json":
            try:
                stocks.append(json.loads(f.read_text(encoding="utf-8")))
            except Exception:
                pass
    return stocks
DOCUMENTS_DIR = Path.home() / "Documents" / "Agentest_Reports"

HEADER_FILL = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)

ACTION_FILLS = {
    "BUY":  PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
    "SELL": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
    "HOLD": PatternFill(start_color="FFD966", end_color="FFD966", fill_type="solid"),
    "WATCH": PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid"),
    "SKIP": PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid"),
}
MATCH_FILL = PatternFill(start_color="006400", end_color="006400", fill_type="solid")
MATCH_FONT = Font(bold=True, color="FFFFFF", size=10)

VERDICT_FILLS = {
    "buy":  PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
    "sell": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
    "hold": PatternFill(start_color="FFD966", end_color="FFD966", fill_type="solid"),
    "watch": PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid"),
}


def _write_cell(ws, row, col, value, font=None, fill=None):
    cell = ws.cell(row=row, column=col, value=value)
    if font:
        cell.font = font
    if fill:
        cell.fill = fill
    cell.border = THIN_BORDER
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    return cell


def _write_header(ws, row, headers):
    for col, h in enumerate(headers, 1):
        _write_cell(ws, row, col, h, font=HEADER_FONT, fill=HEADER_FILL)
    ws.row_dimensions[row].height = 24


def generate_enriched_report(analyze_fundamentals: bool = True, folder: Path | None = None) -> Path:
    """Generate Excel report with technical signals + optional fundamental analysis."""
    signals = _load_all_results()
    if not signals:
        msg = "No scan results found in tests/temp/scan_results/"
        print(f"[!] {msg}")
        raise FileNotFoundError(msg)

    # Run fundamental analysis if requested
    fundamental_data: list[dict] = []
    if analyze_fundamentals:
        buy_sell = [s for s in signals if s["action"] in ("BUY", "SELL")]
        if buy_sell:
            print(f"[+] Running AI fundamental analysis on {len(buy_sell)} signals...")
            fundamental_data = analyze_batch(buy_sell)
        else:
            print("[+] No BUY/SELL signals to analyze fundamentally")

    # Build lookup
    fa_lookup: dict[str, dict] = {}
    for fa in fundamental_data:
        fa_lookup[fa["symbol"]] = fa

    ts = datetime.now()
    folder = folder or DOCUMENTS_DIR / ts.strftime("%Y%m%d_%H%M%S")
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"FocusScan_{ts.strftime('%Y%m%d_%H%M%S')}.xlsx"

    wb = Workbook()

    # ── Summary Sheet (first) ──
    ws = wb.active
    ws.title = "Summary"
    ws.cell(row=1, column=1, value="Focus Scan Report").font = Font(bold=True, size=14)
    ws.cell(row=2, column=1, value=f"Generated: {ts:%Y-%m-%d %H:%M:%S}").font = Font(size=10, italic=True)

    counts: dict[str, int] = {}
    for s in signals:
        a = s.get("action", "SKIP")
        counts[a] = counts.get(a, 0) + 1

    row = 4
    summary_headers = ["Action", "Count"]
    for col, h in enumerate(summary_headers, 1):
        _write_cell(ws, row, col, h, font=HEADER_FONT, fill=HEADER_FILL)
    row += 1
    for action in ["BUY", "SELL", "HOLD", "WATCH", "SKIP"]:
        c = counts.get(action, 0)
        fill = ACTION_FILLS.get(action, PatternFill())
        _write_cell(ws, row, 1, action, fill=fill)
        _write_cell(ws, row, 2, c, fill=fill)
        row += 1

    # AI verdict summary
    if fundamental_data:
        row += 1
        _write_cell(ws, row, 1, "AI Verdict Summary", font=Font(bold=True, size=12))
        row += 1
        v_headers = ["Verdict", "Count"]
        for col, h in enumerate(v_headers, 1):
            _write_cell(ws, row, col, h, font=HEADER_FONT, fill=HEADER_FILL)
        row += 1
        v_counts: dict[str, int] = {}
        for fa in fundamental_data:
            v = fa.get("overall_verdict", "hold")
            v_counts[v] = v_counts.get(v, 0) + 1
        for verdict in ["buy", "sell", "hold", "watch"]:
            c = v_counts.get(verdict, 0)
            fill = VERDICT_FILLS.get(verdict, PatternFill())
            _write_cell(ws, row, 1, verdict.upper(), fill=fill)
            _write_cell(ws, row, 2, c, fill=fill)
            row += 1

    # BUY / SELL symbol lists split by category with AI verdict + match highlight
    row += 2
    _write_cell(ws, row, 1, "Recommended Symbols", font=Font(bold=True, size=12))
    row += 1
    for action in ("BUY", "SELL"):
        etfs = sorted([s for s in signals if s.get("action") == action and s.get("category") == "etf"], key=lambda x: x.get("symbol", ""))
        stocks = sorted([s for s in signals if s.get("action") == action and s.get("category") == "large_cap"], key=lambda x: x.get("symbol", ""))
        if not etfs and not stocks:
            continue
        expected_verdict = action.lower()
        _write_cell(ws, row, 1, f"{action} ({len(etfs)+len(stocks)})", fill=ACTION_FILLS.get(action))
        row += 1
        _write_cell(ws, row, 1, "ETFs", font=Font(bold=True))
        _write_cell(ws, row, 2, "Verdict", font=Font(bold=True))
        _write_cell(ws, row, 3, "Stocks", font=Font(bold=True))
        _write_cell(ws, row, 4, "Verdict", font=Font(bold=True))
        row += 1
        max_len = max(len(etfs), len(stocks), 1)
        for i in range(max_len):
            match_any = False
            col_vals = []
            for col_idx, items in enumerate([etfs, stocks]):
                sym_data = items[i] if i < len(items) else None
                if sym_data:
                    sym = sym_data.get("symbol", "")
                    fa = fa_lookup.get(sym, {})
                    verdict = fa.get("overall_verdict", "")
                    is_match = verdict == expected_verdict
                    if is_match:
                        match_any = True
                    col_vals.append((sym, verdict, is_match))
                else:
                    col_vals.append(("", "", False))
            fill = MATCH_FILL if match_any else None
            font = MATCH_FONT if match_any else None
            c1, v1, m1 = col_vals[0]
            c2, v2, m2 = col_vals[1]
            _write_cell(ws, row, 1, c1, fill=fill if m1 else None, font=font if m1 else None)
            _write_cell(ws, row, 2, v1, fill=fill if m1 else None, font=font if m1 else None)
            _write_cell(ws, row, 3, c2, fill=fill if m2 else None, font=font if m2 else None)
            _write_cell(ws, row, 4, v2, fill=fill if m2 else None, font=font if m2 else None)
            row += 1
        row += 1

    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 26
    ws.column_dimensions["D"].width = 12

    # ── All Signals Sheet ──
    ws2 = wb.create_sheet("All Signals")
    headers = [
        "Symbol", "Category", "Action", "Reason", "Priority",
        "Close", "EMA(10)", "EMA(20)", "Crossover", "HA Signal", "HA Strength",
        "Fundamental Sentiment", "Sentiment Score",
        "Fundamental Support", "Fundamental Details",
        "AI Verdict", "Verdict Reason", "Key News",
    ]
    _write_header(ws2, 1, headers)

    for i, s in enumerate(signals, 2):
        fa = fa_lookup.get(s["symbol"], {})
        values = [
            s.get("symbol", ""), s.get("category", ""), s.get("action", ""),
            s.get("reason", ""), s.get("priority", ""), s.get("close", ""),
            s.get("ema_10", ""), s.get("ema_20", ""), s.get("crossover", ""),
            s.get("ha_signal", ""), s.get("ha_strength", ""),
            fa.get("sentiment", ""), fa.get("sentiment_score", ""),
            fa.get("fundamental_support", ""), fa.get("fundamental_details", ""),
            fa.get("overall_verdict", ""), fa.get("verdict_reason", ""),
            "\n".join(fa.get("key_news", [])),
        ]
        action = s.get("action", "SKIP")
        fill = ACTION_FILLS.get(action, PatternFill())
        for col, val in enumerate(values, 1):
            cell = ws2.cell(row=i, column=col, value=val)
            cell.fill = fill
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws2.row_dimensions[i].height = 36

    widths = [16, 14, 10, 32, 10, 10, 10, 10, 14, 14, 12,
              16, 12, 14, 36, 10, 36, 40]
    for col, w in enumerate(widths, 1):
        ws2.column_dimensions[get_column_letter(col)].width = w
    ws2.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(signals)+1}"
    ws2.freeze_panes = "A2"

    # ── AI Verdicts Sheet ──
    if fundamental_data:
        ws3 = wb.create_sheet("AI Analysis")
        fa_headers = [
            "Symbol", "Category", "Technical Action",
            "Sentiment", "Score", "Fundamental Support",
            "AI Verdict", "Verdict Reason", "Key News",
        ]
        _write_header(ws3, 1, fa_headers)
        for i, fa in enumerate(fundamental_data, 2):
            verdict = fa.get("overall_verdict", "")
            fill = VERDICT_FILLS.get(verdict, PatternFill())
            values = [
                fa.get("symbol", ""), fa.get("category", ""),
                fa.get("technical_action", ""), fa.get("sentiment", ""),
                fa.get("sentiment_score", ""), fa.get("fundamental_support", ""),
                verdict, fa.get("verdict_reason", ""),
                "\n".join(fa.get("key_news", [])),
            ]
            for col, val in enumerate(values, 1):
                cell = ws3.cell(row=i, column=col, value=val)
                cell.fill = fill
                cell.border = THIN_BORDER
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            ws3.row_dimensions[i].height = 36

        fa_widths = [16, 14, 14, 14, 10, 16, 10, 36, 40]
        for col, w in enumerate(fa_widths, 1):
            ws3.column_dimensions[get_column_letter(col)].width = w

    wb.save(path)
    return path
