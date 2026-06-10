from __future__ import annotations
import json
from datetime import datetime
from pathlib import Path
from typing import Any
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SCAN_RESULTS_DIR = Path("tests/temp/scan_results")
DOCUMENTS_DIR = Path.home() / "Documents" / "Agentest_Reports"

HEADERS = [
    ("Symbol", 14), ("Name", 30), ("Category", 12), ("Close", 10),
    ("EMA(10)", 10), ("EMA(20)", 10), ("Crossover", 16),
    ("HA Signal", 14), ("HA Strength", 12),
    ("Action", 10), ("Reason", 40), ("Priority", 10),
]

ACTION_COLORS = {
    "BUY":  PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
    "SELL": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
    "HOLD": PatternFill(start_color="FFD966", end_color="FFD966", fill_type="solid"),
    "WATCH": PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid"),
    "SKIP": PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid"),
}

ACTION_FONTS = {
    "BUY":  Font(bold=True, color="006100"),
    "SELL": Font(bold=True, color="9C0006"),
    "HOLD": Font(bold=True, color="9C6500"),
    "WATCH": Font(bold=True, color="595959"),
    "SKIP": Font(bold=True, color="808080"),
}

HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)


def _load_results() -> dict[str, list[dict[str, Any]]]:
    by_category: dict[str, list[dict]] = {}
    if not SCAN_RESULTS_DIR.exists():
        return by_category
    for f in sorted(SCAN_RESULTS_DIR.iterdir()):
        if f.suffix != ".json":
            continue
        try:
            d = json.loads(f.read_text(encoding="utf-8"))
            cat = d.get("category", "unknown")
            by_category.setdefault(cat, []).append(d)
        except Exception:
            pass
    return by_category


def _write_header(ws, row: int = 1) -> int:
    for col, (label, width) in enumerate(HEADERS, 1):
        cell = ws.cell(row=row, column=col, value=label)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[row].height = 22
    return row + 1


def _write_row(ws, row: int, d: dict) -> int:
    name = d.get("name", d["symbol"]) if "name" in d else d["symbol"]
    values = [
        d.get("symbol", ""), name, d.get("category", ""),
        d.get("close", ""), d.get("ema_10", ""), d.get("ema_20", ""),
        d.get("crossover", ""), d.get("ha_signal", ""), d.get("ha_strength", ""),
        d.get("action", ""), d.get("reason", ""), d.get("priority", ""),
    ]
    action = d.get("action", "SKIP")
    fill = ACTION_COLORS.get(action, PatternFill())
    font = ACTION_FONTS.get(action, Font())
    for col, val in enumerate(values, 1):
        cell = ws.cell(row=row, column=col, value=val)
        cell.fill = fill
        cell.font = font
        cell.border = THIN_BORDER
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 20
    return row + 1


def _write_category_sheet(wb: Workbook, cat: str, stocks: list[dict]) -> None:
    ws = wb.create_sheet(title=cat)
    row = _write_header(ws)
    for d in stocks:
        row = _write_row(ws, row, d)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{row - 1}"
    ws.freeze_panes = "A2"


def _write_summary_sheet(wb: Workbook, by_category: dict[str, list[dict]]) -> None:
    ws = wb.active
    ws.title = "Summary"

    ws.cell(row=1, column=1, value="Daily Scan Report").font = Font(bold=True, size=14)
    ws.cell(row=2, column=1, value=f"Generated: {datetime.now():%Y-%m-%d %H:%M:%S}").font = Font(size=10, italic=True)
    ws.merge_cells("A1:E1")

    row = 4
    headers = ["Category", "Total", "BUY", "SELL", "HOLD", "WATCH", "SKIP"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER
        cell.alignment = Alignment(horizontal="center")
    ws.column_dimensions["A"].width = 16
    for c in range(2, 8):
        ws.column_dimensions[get_column_letter(c)].width = 10
    row += 1

    grand = {k: 0 for k in ["BUY", "SELL", "HOLD", "WATCH", "SKIP"]}
    grand_total = 0
    for cat in sorted(by_category):
        stocks = by_category[cat]
        counts = {k: sum(1 for s in stocks if s.get("action") == k) for k in grand}
        total = len(stocks)
        grand_total += total
        for k in grand:
            grand[k] += counts[k]
        vals = [cat, total] + [counts[k] for k in grand]
        for col, v in enumerate(vals, 1):
            cell = ws.cell(row=row, column=col, value=v)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="center")
            if col >= 3:
                action_key = list(grand.keys())[col - 3]
                cell.fill = ACTION_COLORS.get(action_key, PatternFill())
                cell.font = ACTION_FONTS.get(action_key, Font())
        row += 1

    ws.cell(row=row, column=1, value="GRAND TOTAL").font = Font(bold=True)
    ws.cell(row=row, column=2, value=grand_total).font = Font(bold=True)
    for col, k in enumerate(grand, 3):
        cell = ws.cell(row=row, column=col, value=grand[k])
        cell.font = Font(bold=True)
        cell.fill = ACTION_COLORS.get(k, PatternFill())
        cell.border = THIN_BORDER
        cell.alignment = Alignment(horizontal="center")


def generate_excel_report() -> Path:
    by_category = _load_results()
    ts = datetime.now()
    folder = DOCUMENTS_DIR / ts.strftime("%Y%m%d_%H%M%S")
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"DailyScan_{ts.strftime('%Y%m%d_%H%M%S')}.xlsx"

    wb = Workbook()
    _write_summary_sheet(wb, by_category)
    for cat in sorted(by_category):
        _write_category_sheet(wb, cat, by_category[cat])
    wb.save(path)
    return path
